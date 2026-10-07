"""خواندن اطلاعات فایل اکسل (xlsx/xlsm) با openpyxl."""

from __future__ import annotations

import datetime as _dt
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


@dataclass
class ColumnInfo:
    index: int          # شماره‌ی ستون (۱ برای A)
    letter: str         # نام ستون (A، B، …)
    title: str          # عنوان سرستون
    unique_title: str   # عنوان یکتا (در صورت تکرار، شماره‌دار)


@dataclass
class SheetInfo:
    name: str
    index: int
    columns: list[ColumnInfo] = field(default_factory=list)
    rows: list[list] = field(default_factory=list)
    header_row: int = 1

    @property
    def row_count(self) -> int:
        return len(self.rows)

    def cell(self, row_index: int, column_index: int) -> str:
        """مقدار یک خانه به‌صورت متن (row_index از ۰)."""
        if not (0 <= row_index < len(self.rows)):
            return ""
        row = self.rows[row_index]
        if not (1 <= column_index <= len(row)):
            return ""
        return cell_text(row[column_index - 1])

    def sample_rows(self, count: int = 5) -> list[list[str]]:
        out = []
        for row in self.rows[:count]:
            out.append([cell_text(v) for v in row[: len(self.columns)]])
        return out


def cell_text(value) -> str:
    """تبدیل مقدار سلول به متن تمیز."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "بله" if value else "خیر"
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return ("%f" % value).rstrip("0").rstrip(".")
    if isinstance(value, int):
        return str(value)
    if isinstance(value, (_dt.datetime, _dt.date, _dt.time)):
        if isinstance(value, _dt.datetime):
            return value.strftime("%Y-%m-%d %H:%M") if (value.hour or value.minute) else value.strftime("%Y-%m-%d")
        if isinstance(value, _dt.time):
            return value.strftime("%H:%M")
        return value.strftime("%Y-%m-%d")
    if isinstance(value, _dt.timedelta):
        total = int(value.total_seconds())
        return f"{total // 3600:02d}:{total % 3600 // 60:02d}:{total % 60:02d}"
    return str(value).strip()


def _unique_titles(titles: list[str]) -> list[str]:
    seen: dict[str, int] = {}
    out: list[str] = []
    for title in titles:
        base = title or "بدون‌نام"
        count = seen.get(base, 0) + 1
        seen[base] = count
        out.append(base if count == 1 else f"{base} ({count})")
    return out


def read_sheets(path: str | Path, header_row: int = 1) -> list[SheetInfo]:
    """همه‌ی شیت‌های فایل اکسل را می‌خواند."""
    wb = load_workbook(filename=str(path), data_only=True, read_only=False)
    sheets: list[SheetInfo] = []
    for index, ws in enumerate(wb.worksheets):
        rows_iter = ws.iter_rows(values_only=True)
        all_rows = [list(r) for r in rows_iter]
        if not all_rows:
            sheets.append(SheetInfo(name=ws.title, index=index, header_row=header_row))
            continue

        guard = min(max(header_row - 1, 0), len(all_rows) - 1)
        header = all_rows[guard]
        width = len(header)
        # عرض واقعی جدول = بیشترین ستون پرِ داده‌ها
        for row in all_rows[guard + 1 :]:
            trimmed = len(row)
            if trimmed > width:
                width = trimmed
        while width > 0 and not cell_text(header[width - 1]) and all(
            not cell_text(r[width - 1]) if len(r) >= width else True for r in all_rows[guard + 1 :]
        ):
            width -= 1

        titles = [cell_text(h) for h in header[:width]]
        unique = _unique_titles(titles)
        columns = [
            ColumnInfo(index=i + 1, letter=get_column_letter(i + 1), title=titles[i], unique_title=unique[i])
            for i in range(width)
        ]

        data_rows = []
        for row in all_rows[guard + 1 :]:
            values = list(row[:width]) + [None] * max(0, width - len(row))
            if all(cell_text(v) == "" for v in values):
                continue
            data_rows.append(values)

        sheets.append(
            SheetInfo(name=ws.title, index=index, columns=columns, rows=data_rows, header_row=header_row)
        )
    return sheets


def resolve_column(sheet: SheetInfo, key) -> ColumnInfo | None:
    """پیدا کردن ستون بر اساس حرف، شماره یا عنوان."""
    if key is None or key == "":
        return None
    if isinstance(key, int):
        for column in sheet.columns:
            if column.index == key:
                return column
        return None
    text = str(key).strip()
    if text.isdigit():
        number = int(text)
        for column in sheet.columns:
            if column.index == number:
                return column
        return text_index(sheet, number - 1)
    for column in sheet.columns:
        if column.letter.lower() == text.lower():
            return column
    for column in sheet.columns:
        if column.title.strip() == text or column.unique_title.strip() == text:
            return column
    return None


def text_index(sheet: SheetInfo, index: int) -> ColumnInfo | None:
    if 0 <= index < len(sheet.columns):
        return sheet.columns[index]
    return None


def normalize_title(value: str) -> str:
    """نرمال‌سازی عنوان ستون/توکن برای مقایسه‌ی هوشمند."""
    text = (value or "").strip().lower()
    for ch in ["\u200c", "‌", "_", "-", ".", "(", ")", "[", "]", "،", ",", "«", "»", "‏", "‎"]:
        text = text.replace(ch, " ")
    text = text.replace("ي", "ی").replace("ك", "ک")
    text = text.replace("أ", "ا").replace("إ", "ا").replace("آ", "ا").replace("ة", "ه")
    return " ".join(text.split())
